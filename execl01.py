import docx
import re,os
import openpyxl

#将docx文件转换成text格式
ducoment = docx.Document(r'test.docx')
for i in ducoment.paragraphs:
    flie_te = i.text
    with open('file.text','a',encoding='utf-8') as f:
        f.write(flie_te)
        f.write('\r\n')

#去掉文件中的空行
file1 = open('file.text','r',encoding='utf-8')
file2 = open('file_now','w',encoding='utf-8')
for lin in file1:
    if  lin == '\n':
        lin = lin.strip('\n')
    else:
        file2.write(lin)
file1.close()
file2.close()
os.remove('file.text')

#获取文档数据，并分类追加到列表中
rule = []
rule_1 = []
rule_2 = []
rule_3 = []
with open('file_now','r',encoding='utf-8') as rf:
    for r in rf:
        res = re.findall('(.*?。)+',r)
        if res != []:
            rule.append(res)
        res1 = re.findall('[A][\.](?:.*)\W',r)
        if res1 != []:
            rule_1.append(res1)
        res2 = re.findall('[（][A-Z]+[）]',r)
        if res2 != []:
            rule_2.append(res2)
        res3 = re.findall('[（][对|错][）]',r)
        if res3 != []:
            rule_3.append(res3)

#将列表中的数据依次添加到execl表格中
wb = openpyxl.Workbook()
ws = wb.active
def ins_data(v1,v2):
    count = 1
    for j in v1:
        ws.cell(row=count,column=v2).value=j[0]
        count +=1

if __name__=="__main__":
    ins_data(rule,1)
    ins_data(rule_1,5)
    ins_data(rule_2,10)
    ins_data(rule_3,15)

wb.save('test.xlsx')













               



